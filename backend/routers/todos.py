from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from auth import get_current_user
import models
import schemas
from typing import List
import json
import io

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(
    prefix="/todos",
    tags=["Todos"]
)

# Create Todo :---

@router.post("/", response_model=schemas.TodoResponse)
def create_todo(todo: schemas.TodoCreate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):

    # Subtasks ko json string mein convert karo
    subtasks_json = json.dumps([s.dict() for s in todo.subtasks])

    new_todo = models.Todo(
        title=todo.title,
        priority=todo.priority,
        tags=todo.tags,
        deadline_date=todo.deadline_date,
        deadline_time=todo.deadline_time,
        subtasks=subtasks_json,
        user_id=current_user.id
    )
    db.add(new_todo)
    db.commit()
    db.refresh(new_todo)

    # Json string wapas list mein convert karo response ke liye
    new_todo.subtasks = json.loads(new_todo.subtasks or "[]")
    
    return new_todo

# Get All Todos :---

@router.get("/", response_model=List[schemas.TodoResponse])
def get_todos(db: Session = Depends(get_db), 
              current_user: models.User = Depends(get_current_user)):
    
    todos = db.query(models.Todo).filter(
        models.Todo.user_id == current_user.id).all()
    
    # Har todo ke subtasks ko JSON string se list mein convert karo
    for todo in todos:
        todo.subtasks = json.loads(todo.subtasks or "[]")

    return todos

# Export Todos as PDF :---

@router.get("/export/pdf")
def export_todos_pdf(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    todos = db.query(models.Todo).filter(models.Todo.user_id == current_user.id).all()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5*cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Smart Todo — Task Export", styles['Title']))
    elements.append(Spacer(1, 12))

    priority_map = {1: "Low", 2: "Medium", 3: "High"}
    data = [["Title", "Priority", "Tags", "Deadline", "Status"]]

    for t in todos:
        deadline_str = f"{t.deadline_date or ''} {str(t.deadline_time or '')[:5]}"
        data.append([
            t.title,
            priority_map.get(t.priority, str(t.priority)),
            t.tags or "-",
            deadline_str.strip() or "-",
            "Completed" if t.completed else "Pending",
        ])

    table = Table(data, repeatRows=1, colWidths=[5.5*cm, 2.5*cm, 3*cm, 3.5*cm, 2.5*cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f6f9')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=smart_todo_export.pdf"}
    )


# Export Todos as Excel :---

@router.get("/export/excel")
def export_todos_excel(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    todos = db.query(models.Todo).filter(models.Todo.user_id == current_user.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Todos"

    priority_map = {1: "Low", 2: "Medium", 3: "High"}
    headers = ["Title", "Priority", "Tags", "Deadline Date", "Deadline Time", "Status"]
    ws.append(headers)

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for t in todos:
        ws.append([
            t.title,
            priority_map.get(t.priority, str(t.priority)),
            t.tags or "-",
            str(t.deadline_date or "-"),
            str(t.deadline_time or "-")[:5],
            "Completed" if t.completed else "Pending",
        ])

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=smart_todo_export.xlsx"}
    )


# Get Single Todo :---

@router.get("/{todo_id}", response_model=schemas.TodoResponse)
def get_single_todo(
    todo_id: int, 
    db: Session = Depends(get_db), 
    current_user: models.User = Depends(get_current_user)):

    # todo = get_todo_or_404(todo_id, current_user, db)
    todo = db.query(models.Todo).filter(
        models.Todo.id == todo_id,
        models.Todo.user_id == current_user.id
    ).first()
    if not todo:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Todo not found"
        )
    todo.subtasks = json.loads(todo.subtasks or "[]")
    return todo

# Update Todo :---

@router.put("/{todo_id}", response_model=schemas.TodoResponse)
def update_todo(
    todo_id: int, 
    updated_todo: schemas.TodoUpdate, 
    db: Session = Depends(get_db), 
    current_user: models.User = Depends(get_current_user)):

    # todo = get_todo_or_404(todo_id, current_user, db)

    todo = db.query(models.Todo).filter(
        models.Todo.id == todo_id,
        models.Todo.user_id == current_user.id
    ).first()
    if not todo:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Todo not found"
        )
    todo.title = updated_todo.title
    todo.completed = updated_todo.completed
    todo.priority = updated_todo.priority
    todo.tags = updated_todo.tags
    todo.deadline_date = updated_todo.deadline_date
    todo.deadline_time = updated_todo.deadline_time
    todo.subtasks = json.dumps([s.dict() for s in updated_todo.subtasks])  

    db.commit()
    db.refresh(todo)
    todo.subtasks = json.loads(todo.subtasks or "[]")
    return todo

# Mark Complete :---

@router.patch("/{todo_id}/complete", response_model=schemas.TodoResponse)
def mark_complete(
    todo_id: int, 
    db: Session = Depends(get_db), 
    current_user: models.User = Depends(get_current_user)):

    # todo = get_todo_or_404(todo_id, current_user, db)

    todo = db.query(models.Todo).filter(
        models.Todo.id == todo_id,
        models.Todo.user_id == current_user.id
    ).first()
    if not todo:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Todo not found"
        )
    todo.completed = not todo.completed

    db.commit()
    db.refresh(todo)
    todo.subtasks = json.loads(todo.subtasks or "[]")
    
    return todo

# Delete Todo :---

@router.delete("/{todo_id}")
def delete_todo(todo_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    todo = db.query(models.Todo).filter(
        models.Todo.id == todo_id,
        models.Todo.user_id == current_user.id
    ).first()
    if not todo:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Todo not found"
        )
    db.delete(todo)
    db.commit()
    return {"message": "Todo deleted successfully"}